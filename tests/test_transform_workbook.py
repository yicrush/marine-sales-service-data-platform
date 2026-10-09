"""Workbook/CLI tests use synthetic documents, never private uploaded files."""

from contextlib import redirect_stdout, redirect_stderr
import hashlib
from io import StringIO
import json
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from openpyxl import Workbook

from etl.transform.__main__ import main
from etl.transform.workbook import transform_workbook


def sample_workbook(path):
    workbook = Workbook()
    m = workbook.active
    m.title = "M sample"
    for coordinate, value in {
        "A1": "REF. NO. KP-M-Q-26-001", "A2": "TO : Sample Customer",
        "A3": "DATE : 15.Jan.2026", "A13": "Sample quotation",
        "B15": "No.", "C15": "Item", "D15": "Description",
        "H15": "Qty", "I15": "U/Price(USD)", "J15": "Amount(USD)",
        "B16": 1, "C16": "Seal", "H16": 2, "I16": 10, "J16": 20,
        "H18": "DISCOUNT", "I18": 0.1, "J18": 2,
        "H19": "TOTAL", "J19": 18,
    }.items():
        m[coordinate] = value
    m["I18"].number_format = "0%"
    kit = workbook.create_sheet("KIT sample")
    kit.append(["Sample KIT"])
    kit.append(["No.", "Parts name", "Part No.", "Quantity", None, "U/Price(USD)", "Amount(USD)"])
    kit.append([None, "Main pump (MODEL-1: 2 sets)", None, None, None, 20])
    kit.append([1, "Seal", "001", 2, "pcs", 10, 20])
    kit.append(["AMOUNT", None, None, None, None, 20])
    cost = workbook.copy_worksheet(kit)
    cost.title = "Cost sample"
    cost["F2"] = "Unit price(KRW)"
    cost["G2"] = "Amount(KRW)"
    cost["C4"] = "002"
    workbook.create_sheet("Notes")["A1"] = "No quotation on this sheet"
    workbook.save(path)
    workbook.close()


class WorkbookTests(unittest.TestCase):
    def setUp(self):
        self.directory = TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.path = Path(self.directory.name) / "sample.xlsx"
        sample_workbook(self.path)

    def test_read_only_extraction_currency_percent_and_explicit_support(self):
        before = hashlib.sha256(self.path.read_bytes()).hexdigest()
        result = transform_workbook(self.path, sheet_names=["M sample"],
                                    kit_sheet_names=["KIT sample"], cost_sheet_names=["Cost sample"])
        record = result["records"][0]
        self.assertTrue(record["is_valid"], record["issues"])
        self.assertEqual(record["quotation"]["quotation_date"], "2026-01-15")
        self.assertEqual(record["quotation"]["currency"], "USD")
        self.assertEqual(record["quotation"]["document_discount_rate"], 10)
        self.assertEqual(record["quotation"]["document_total_amount"], 18)
        self.assertEqual([d["document"]["currency"] for d in record["kit_documents"]], ["USD", "KRW"])
        self.assertEqual([d["groups"][0]["components"][0]["part_no"] for d in record["kit_documents"]], ["001", "002"])
        self.assertEqual(result["skipped_sheets"], ["Notes"])
        self.assertEqual(hashlib.sha256(self.path.read_bytes()).hexdigest(), before)

    def test_default_ignores_supporting_sheets_and_does_not_guess_links(self):
        result = transform_workbook(self.path)
        self.assertEqual(len(result["records"]), 1)
        self.assertEqual(result["records"][0]["kit_documents"], [])
        self.assertIn("KIT sample", result["skipped_sheets"])

    def test_numberless_priced_row_reaches_load_validation(self):
        from openpyxl import load_workbook
        from etl.load.validation import validate_records
        workbook = load_workbook(self.path)
        workbook["M sample"]["B16"] = None
        workbook.save(self.path)
        workbook.close()
        before = hashlib.sha256(self.path.read_bytes()).hexdigest()
        result = transform_workbook(self.path, sheet_names=["M sample"])
        accepted, rejected = validate_records(result["records"])
        self.assertEqual(rejected, [])
        self.assertEqual(len(accepted), 1)
        item = accepted[0]["quotation_items"][0]
        self.assertEqual((item["line_no"], item["quantity"], item["unit_price"], item["quoted_amount"]), (1, 2, 10, 20))
        self.assertIsNone(result["records"][0]["raw"]["items"][0]["source_line_no"])
        self.assertEqual(result["records"][0]["raw"]["items"][0]["source_row"], 16)
        self.assertEqual(hashlib.sha256(self.path.read_bytes()).hexdigest(), before)

    def test_restarted_numbers_preserve_distinct_section_context(self):
        from openpyxl import load_workbook
        from etl.load.validation import validate_records
        workbook = load_workbook(self.path)
        sheet = workbook["M sample"]
        for row in sheet.iter_rows(min_row=16, max_row=25):
            for cell in row:
                cell.value = None
        for coordinate, value in {
            "B16": "Main pump model A", "B17": 1, "C17": "Seal", "H17": 2, "I17": 10, "J17": 20,
            "B19": "Main pump model B", "B20": 1, "C20": "Seal", "H20": 3, "I20": 10, "J20": 30,
            "H22": "TOTAL(USD)", "J22": 50,
            "A24": "TERMS AND CONDITIONS:", "B25": 1, "C25": "Unrelated footer", "H25": 1, "I25": 999, "J25": 999,
        }.items():
            sheet[coordinate] = value
        workbook.save(self.path)
        workbook.close()
        result = transform_workbook(self.path, sheet_names=["M sample"])
        accepted, rejected = validate_records(result["records"])
        self.assertEqual(rejected, [])
        items = accepted[0]["quotation_items"]
        self.assertEqual([item["line_no"] for item in items], [1, 2])
        self.assertEqual([item["quantity"] for item in items], [2, 3])
        self.assertEqual([item["quoted_amount"] for item in items], [20, 30])
        self.assertIn("SECTION: Main pump model A", items[0]["line_text_raw"])
        self.assertIn("SECTION: Main pump model B", items[1]["line_text_raw"])
        self.assertIn("LINE: 1", items[1]["line_text_raw"])
        raw_items = result["records"][0]["raw"]["items"]
        self.assertEqual([item["source_line_no"] for item in raw_items], [1, 1])
        self.assertEqual([item["source_row"] for item in raw_items], [17, 20])

    def test_per_vessel_displayed_price_is_not_an_aggregate_total(self):
        from etl.transform import transform_quotation
        record = transform_quotation({
            "type": "M", "header": {"ref_no_raw": "REF. NO. KP-M-Q-26-001", "customer_raw": "TO : Sample"},
            "summary": {"quoted_amount": "USD 15,000 / VESSEL"},
        }, currency="USD")
        self.assertIsNone(record["quotation"]["document_total_amount"])
        self.assertEqual(record["pricing_context"]["displayed_amount_per_unit"], 15000)
        self.assertEqual(record["pricing_context"]["amount_basis_raw"], "VESSEL")
        self.assertIn("per_unit_total", [issue["code"] for issue in record["issues"]])

    def test_numeric_total_with_per_vessel_label_requires_review(self):
        from openpyxl import load_workbook
        from etl.load.validation import ValidationError, validate_records
        workbook = load_workbook(self.path)
        workbook["M sample"]["H19"] = "TOTAL / vessel"
        workbook.save(self.path)
        workbook.close()
        record = transform_workbook(self.path, sheet_names=["M sample"])["records"][0]
        self.assertIsNone(record["quotation"]["document_total_amount"])
        self.assertEqual(record["pricing_context"]["displayed_amount_per_unit"], 18)
        self.assertEqual(record["pricing_context"]["amount_basis_raw"], "vessel")
        self.assertEqual(record["raw"]["summary"]["quoted_amount_label_raw"], "TOTAL / vessel")
        with self.assertRaises(ValidationError):
            validate_records([record])
        accepted, rejected = validate_records([record], allow_warnings=True)
        self.assertEqual(rejected, [])
        self.assertIsNone(accepted[0]["quotation"]["document_total_amount"])

    def test_full_named_month_with_dots(self):
        from etl.transform.common import parse_date, ParseError
        self.assertEqual(parse_date("DATE : 22.July.2025"), "2025-07-22")
        self.assertEqual(parse_date("DATE : 30.June.2025"), "2025-06-30")
        with self.assertRaises(ParseError):
            parse_date("DATE : 07.Junly.2026")

    def test_worker_count_after_worker_label(self):
        from etl.transform import transform_quotation
        record = transform_quotation({
            "type": "O", "header": {"ref_no_raw": "KP-O-Q-26-001", "customer_raw": "Sample"},
            "service_details": {"service_engineers_raw": "Our Two Local S/E and customer to provide shipyard worker 2-person - 9 days"},
        }, currency="USD")
        self.assertEqual(record["service_details"]["service_engineer_count"], 2)
        self.assertEqual(record["service_details"]["supporting_worker_count"], 2)
        self.assertEqual(record["service_details"]["working_days"], 9)

    def test_invalid_selection_and_double_support_type_fail(self):
        for kwargs in ({"sheet_names": ["Missing"]}, {"sheet_names": ["Notes"]},
                       {"sheet_names": []}, {"sheet_names": ["M sample"],
                        "kit_sheet_names": ["KIT sample"], "cost_sheet_names": ["KIT sample"]}):
            with self.subTest(kwargs=kwargs), self.assertRaises(ValueError):
                transform_workbook(self.path, **kwargs)

    def test_uncached_formula_is_reported_without_calculating(self):
        from openpyxl import load_workbook
        workbook = load_workbook(self.path)
        workbook["M sample"]["J16"] = "=H16*I16"
        workbook.save(self.path)
        workbook.close()
        result = transform_workbook(self.path, sheet_names=["M sample"])
        record = result["records"][0]
        self.assertIsNone(record["quotation_items"][0]["quoted_amount"])
        self.assertIn("uncached_formulas", [issue["code"] for issue in record["issues"]])

    def test_shifted_cost_price_columns_are_errors(self):
        from openpyxl import load_workbook
        workbook = load_workbook(self.path)
        sheet = workbook["Cost sample"]
        sheet["F2"] = None
        sheet["G2"] = "Unit price(KRW)"
        sheet["H2"] = "Amount(KRW)"
        workbook.save(self.path)
        workbook.close()
        result = transform_workbook(self.path, sheet_names=["M sample"], cost_sheet_names=["Cost sample"])
        self.assertFalse(result["records"][0]["is_valid"])
        self.assertIn("shifted_extract_columns", [issue["code"] for issue in result["records"][0]["issues"]])

    def test_cli_writes_decimal_strings_and_preserves_existing_outputs(self):
        output = Path(self.directory.name) / "result.json"
        args = [str(self.path), "--sheet", "M sample", "--output", str(output), "--strict"]
        with redirect_stdout(StringIO()):
            self.assertEqual(main(args), 0)
        payload = json.loads(output.read_text(encoding="utf-8"))
        self.assertEqual(payload["records"][0]["quotation"]["document_total_amount"], "18")
        original = output.read_bytes()
        with redirect_stderr(StringIO()), self.assertRaises(SystemExit) as raised:
            main(args)
        self.assertEqual(raised.exception.code, 2)
        self.assertEqual(output.read_bytes(), original)

    def test_cli_strict_error_still_writes_report(self):
        from openpyxl import load_workbook
        workbook = load_workbook(self.path)
        workbook["M sample"]["H16"] = 1.5
        workbook.save(self.path)
        workbook.close()
        output = Path(self.directory.name) / "invalid.json"
        with redirect_stdout(StringIO()):
            status = main([str(self.path), "--sheet", "M sample", "--output", str(output), "--strict"])
        self.assertEqual(status, 1)
        self.assertFalse(json.loads(output.read_text())["records"][0]["is_valid"])


if __name__ == "__main__":
    unittest.main()
