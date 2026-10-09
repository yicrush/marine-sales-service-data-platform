"""Real PostgreSQL integration tests using only disposable, uniquely named schemas.

Set MARINE_TEST_DATABASE_URL to an explicitly selected development/test database.
No external credentials, confidential Excel files, or public-schema tables are used.
"""

from concurrent.futures import ThreadPoolExecutor
from contextlib import redirect_stderr, redirect_stdout
from copy import deepcopy
from decimal import Decimal
import io
import json
import os
from pathlib import Path
from tempfile import TemporaryDirectory
from threading import Barrier
import unittest
from uuid import uuid4

import psycopg
from psycopg import IsolationLevel, sql
from psycopg.pq import TransactionStatus
from psycopg.rows import dict_row
from openpyxl import Workbook

from etl.load import LoadError, ValidationError, initialize_schema, load_records
from etl.load.__main__ import main as json_load_main
from etl.pipeline import main as pipeline_main
from etl.transform import to_jsonable, transform_kit_document, transform_quotation


TABLES = (
    "customers", "contacts", "vessels", "quotations", "quotation_vessels",
    "quotation_items", "kit_documents", "kit_component_groups", "kit_components",
    "service_details",
)


def _document(kind):
    cost = kind == "COST_SHEET"
    amount, price = ("900.00", "450.00") if cost else ("6.50", "3.25")
    extracted = {
        "document_type": kind,
        "sheet_name": "synthetic cost" if cost else "synthetic kit",
        "title_raw": "Synthetic overhaul KIT\nEXAMPLE VESSEL",
        "amount_summary": {"label_raw": "Amount", "value_raw": amount},
        "groups": [{
            "group_name_raw": "For No.1 Main pump (MODEL-001: 1 set)",
            "group_amount": None,
            "components": [{
                "line_no": 1, "part_name": "O-Ring", "part_no": "000019-A",
                "quantity": 2, "quantity_unit": "pcs", "unit_price": price,
                "amount": amount, "remark": None,
            }],
        }],
    }
    result = transform_kit_document(extracted, source_file_name="synthetic.xlsx",
                                    currency="KRW" if cost else "USD")
    if result["issues"]:
        raise AssertionError(f"Integration document fixture has issues: {result['issues']}")
    return result


def _record(sequence=1, *, kind="M", customer="Example Marine", contact=True,
            documents=False, unknown_total=False):
    """Use the public Transform APIs to create synthetic, warning-free records."""
    extracted = {
        "type": kind,
        "header": {
            "sheet_name": f"synthetic quote {sequence}",
            "customer_raw": f"TO : {customer}",
            "contact_raw": "ATTN : A. Contact" if contact else None,
            "contact_details_raw": "TEL : 001-0200 FAX : 001-0300 E-MAIL : contact@example.test",
            "date_raw": "DATE : 6 January 2026",
            "ref_no_raw": f"REF. NO. : KP-{kind}-Q-26-{sequence:03d}",
            "vessel_name_raw": "VESSEL NAME : EXAMPLE VESSEL",
            "shipyard_hull_raw": "SHIPYARD/HULL NO. : H-001",
            "subject_raw": "Subject : Synthetic quotation",
        },
        "items": [{
            "line_no": 1, "item_raw": "Inspection service" if kind == "O" else "Example part",
            "description_raw": "Original specification\nDo not alter spelling.",
            "part_no": "000019-A", "quantity": 2, "quantity_unit": "pcs",
            "unit_price": "11.70", "amount": "23.40", "remark": None,
        }],
        "summary": {"discount_rate": None, "discounted_amount": None,
                    "quoted_amount": None if unknown_total else "23.40"},
        "terms": {"lead_time_raw": None, "delivery_terms_raw": None,
                  "payment_terms_raw": None},
    }
    if kind == "O":
        extracted["service_details"] = {
            "working_place_raw": "Working place : EXAMPLE PORT",
            "service_engineers_raw": "Number of Service Engineers : 2 service engineers + 1 worker for 8 days including travelling days",
            "working_details_raw": ["A. Inspection", "B. Original scope"],
            "expense_lines_raw": [], "remarks_raw": None,
        }
    supporting = [_document("KIT_DETAIL"), _document("COST_SHEET")] if documents else []
    result = transform_quotation(extracted, source_file_name="synthetic.xlsx",
                                 currency="USD", kit_documents=supporting)
    if result["issues"]:
        raise AssertionError(f"Integration quotation fixture has issues: {result['issues']}")
    return result


def _write_workbook(path, *, valid_customer=True):
    """An invented Excel source compatible with the repository's current Extract."""
    workbook = Workbook()
    quote = workbook.active
    quote.title = "quote"
    cells = {
        "A1": "TO : Example Marine" if valid_customer else None,
        "A2": "ATTN : A. Contact", "A3": "DATE : 6 January 2026",
        "A4": "REF. NO. : KP-M-Q-26-001",
        "A5": "VESSEL NAME : EXAMPLE VESSEL",
        "A6": "SHIPYARD/HULL NO. : H-001", "A7": "TEL : 001-0200",
        "A13": "Subject : Synthetic spares", "B15": "No.", "C15": "ITEM",
        "D15": "Description", "H15": "Qty", "I15": "U/PRICE", "J15": "AMOUNT",
        "B16": 1, "C16": "Example part", "D16": "Original specification",
        "H16": 2, "I16": 11.70, "J16": 23.40,
        "H18": "TOTAL(USD)", "J18": 23.40,
    }
    for address, value in cells.items():
        quote[address] = value
    for name, currency, price, amount in (
        ("kit", "USD", 3.25, 6.50), ("cost", "WON", 450, 900),
    ):
        sheet = workbook.create_sheet(name)
        for address, value in {
            "A1": "Synthetic KIT", "A2": "No.", "B2": "Parts name",
            "C2": "Part no.", "D2": "Qty", "E2": "Unit",
            "F2": f"Unit price({currency})", "G2": f"Amount({currency})",
            "B3": "For No.1 Main pump (MODEL-001: 1 set)",
            "A4": 1, "B4": "O-Ring", "C4": "000019-A", "D4": 2,
            "E4": "pcs", "F4": price, "G4": amount, "A6": "Amount", "F6": amount,
        }.items():
            sheet[address] = value
    try:
        workbook.save(path)
    finally:
        workbook.close()


@unittest.skipUnless(
    os.environ.get("MARINE_TEST_DATABASE_URL"),
    "Set MARINE_TEST_DATABASE_URL to run disposable-schema PostgreSQL integration tests",
)
class PostgreSQLLoadTests(unittest.TestCase):
    def setUp(self):
        self.dsn = os.environ["MARINE_TEST_DATABASE_URL"]
        self.connection = psycopg.connect(self.dsn, autocommit=True, connect_timeout=5)
        self.schema = "marine_test_" + uuid4().hex
        self.addCleanup(self.connection.close)
        self.addCleanup(self._drop_schema)
        initialize_schema(self.connection, schema=self.schema)

    def _drop_schema(self):
        self.connection.execute(sql.SQL("DROP SCHEMA IF EXISTS {} CASCADE").format(
            sql.Identifier(self.schema)))

    def _query(self, statement, parameters=()):
        with self.connection.cursor() as cursor:
            cursor.execute(sql.SQL(statement).format(schema=sql.Identifier(self.schema)), parameters)
            return cursor.fetchall()

    def _count(self, table):
        return self.connection.execute(sql.SQL("SELECT count(*) FROM {}.{}").format(
            sql.Identifier(self.schema), sql.Identifier(table))).fetchone()[0]

    def _counts(self):
        return {table: self._count(table) for table in TABLES}

    def _snapshot(self):
        return {table: self.connection.execute(
            sql.SQL("SELECT * FROM {}.{} ORDER BY 1, 2").format(
                sql.Identifier(self.schema), sql.Identifier(table))).fetchall()
            for table in TABLES}

    def _load(self, records, **kwargs):
        return load_records(self.connection, records, schema=self.schema, **kwargs)

    def _reject_component_trigger(self):
        self.connection.execute(sql.SQL("""
            CREATE FUNCTION {}.reject_component() RETURNS trigger
            LANGUAGE plpgsql AS $body$
            BEGIN
                IF NEW.part_name = 'FAIL_COMPONENT' THEN
                    RAISE EXCEPTION 'synthetic component insertion failure';
                END IF;
                RETURN NEW;
            END
            $body$
        """).format(sql.Identifier(self.schema)))
        self.connection.execute(sql.SQL("""
            CREATE TRIGGER reject_component BEFORE INSERT ON {}.kit_components
            FOR EACH ROW EXECUTE FUNCTION {}.reject_component()
        """).format(sql.Identifier(self.schema), sql.Identifier(self.schema)))

    def _run_cli(self, main, arguments):
        output, errors = io.StringIO(), io.StringIO()
        with redirect_stdout(output), redirect_stderr(errors):
            status = main(arguments)
        combined = output.getvalue() + errors.getvalue()
        self.assertNotIn(self.dsn, combined)
        return status, combined

    def _database_cli_arguments(self, report):
        return ["--database-url-env", "MARINE_TEST_DATABASE_URL", "--schema", self.schema,
                "--report", str(report)]

    def test_full_graph_preserves_foreign_keys_decimals_text_and_nulls(self):
        summary = self._load([
            _record(documents=True), _record(2, kind="O", contact=False, unknown_total=True),
        ])
        self.assertEqual((summary["inserted"], summary["skipped"], summary["rejected"]), (2, 0, 0))
        self.assertEqual(self._counts(), {
            "customers": 1, "contacts": 1, "vessels": 1, "quotations": 2,
            "quotation_vessels": 2, "quotation_items": 2, "kit_documents": 2,
            "kit_component_groups": 2, "kit_components": 2, "service_details": 1,
        })
        rows = self._query("""
            SELECT q.quotation_id, c.customer_id, ct.telephone, v.customer_id,
                   d.document_type, g.equipment_model, k.part_no, k.amount
            FROM {schema}.quotations q
            JOIN {schema}.customers c ON c.customer_id = q.customer_id
            JOIN {schema}.contacts ct ON ct.contact_id = q.contact_id
                                    AND ct.customer_id = c.customer_id
            JOIN {schema}.quotation_vessels qv ON qv.quotation_id = q.quotation_id
            JOIN {schema}.vessels v ON v.vessel_id = qv.vessel_id
            JOIN {schema}.kit_documents d ON d.quotation_id = q.quotation_id
            JOIN {schema}.kit_component_groups g ON g.kit_document_id = d.kit_document_id
            JOIN {schema}.kit_components k ON k.kit_component_group_id = g.kit_component_group_id
            ORDER BY d.document_type
        """)
        self.assertEqual(len(rows), 2)
        for _, customer_id, telephone, vessel_customer, _, model, part_no, amount in rows:
            self.assertEqual(customer_id, vessel_customer)
            self.assertEqual(telephone, "001-0200")
            self.assertEqual(model, "MODEL-001")
            self.assertEqual(part_no, "000019-A")
            self.assertIsInstance(amount, Decimal)
        self.assertEqual([row[-1] for row in rows], [Decimal("900.00"), Decimal("6.50")])
        self.assertEqual(self._query("""
            SELECT contact_id, document_total_amount, document_discount_rate,
                   document_discount_amount FROM {schema}.quotations
            WHERE quotation_type = 'O'
        """), [(None, None, None, None)])
        self.assertEqual(self._query("""
            SELECT part_no, unit_price, quoted_amount, discount_rate
            FROM {schema}.quotation_items ORDER BY quotation_item_id
        """), [("000019-A", Decimal("11.70"), Decimal("23.40"), None)] * 2)
        self.assertEqual(self._query("SELECT group_amount FROM {schema}.kit_component_groups"), [(None,), (None,)])
        self.assertEqual(self._query("""
            SELECT service_engineer_count, supporting_worker_count, working_days,
                   includes_travel_days FROM {schema}.service_details
        """), [(2, 1, 8, True)])

    def test_default_skip_keeps_entire_graph_and_reports_existing_id(self):
        record = _record(documents=True)
        first = self._load([record])
        before = self._snapshot()
        second = self._load([record])
        self.assertEqual((second["inserted"], second["skipped"], second["replaced"]), (0, 1, 0))
        self.assertEqual(second["records"][0]["quotation_id"], first["records"][0]["quotation_id"])
        self.assertEqual(self._snapshot(), before)

    def test_master_reuse_is_exact_and_scoped_to_customer(self):
        self._load([_record(), _record(2)])
        self.assertEqual([self._count(table) for table in ("customers", "contacts", "vessels")], [1, 1, 1])
        self._load([_record(3, customer="example marine")])
        self.assertEqual([self._count(table) for table in ("customers", "contacts", "vessels")], [2, 2, 2])
        self.assertEqual(self._query("SELECT count(DISTINCT contact_id), count(DISTINCT customer_id) FROM {schema}.quotations"), [(2, 2)])

    def test_replace_keeps_quotation_id_removes_owned_children_and_keeps_masters(self):
        original = _record(kind="O", documents=True)
        first = self._load([original])
        old_item_id = self._query("SELECT quotation_item_id FROM {schema}.quotation_items")[0][0]
        replacement = _record(kind="O", contact=False)
        replacement["quotation_items"][0]["description"] = "Replacement specification"
        replacement["vessels"] = []
        replacement["quotation_vessels"] = []
        replacement["service_details"] = None
        result = self._load([replacement], existing_policy="replace")
        self.assertEqual((result["inserted"], result["replaced"]), (0, 1))
        self.assertEqual(result["records"][0]["quotation_id"], first["records"][0]["quotation_id"])
        for table in ("kit_documents", "kit_component_groups", "kit_components", "quotation_vessels", "service_details"):
            with self.subTest(table=table):
                self.assertEqual(self._count(table), 0)
        self.assertEqual([self._count(table) for table in ("customers", "contacts", "vessels")], [1, 1, 1])
        items = self._query("SELECT quotation_item_id, description FROM {schema}.quotation_items")
        self.assertEqual(len(items), 1)
        self.assertNotEqual(items[0][0], old_item_id)
        self.assertEqual(items[0][1], "Replacement specification")
        self.assertEqual(self._query("SELECT contact_id FROM {schema}.quotations"), [(None,)])

    def test_existing_error_rolls_back_earlier_records_and_new_masters(self):
        existing = _record()
        self._load([existing])
        before = self._snapshot()
        with self.assertRaises(LoadError):
            self._load([_record(2, customer="New batch customer"), existing], existing_policy="error")
        self.assertEqual(self._snapshot(), before)

    def test_ambiguous_existing_customer_is_rejected_without_partial_writes(self):
        for _ in range(2):
            self.connection.execute(sql.SQL("INSERT INTO {}.customers (customer_name) VALUES (%s)").format(
                sql.Identifier(self.schema)), ("Example Marine",))
        before = self._snapshot()
        with self.assertRaisesRegex(LoadError, "Ambiguous existing customers"):
            self._load([_record()])
        self.assertEqual(self._snapshot(), before)

    def test_source_sql_injection_text_is_stored_literally(self):
        source = "x'); DROP TABLE quotations; --\nnot executable.xlsx"
        record = _record(customer="Example'); DELETE FROM customers; --")
        record["quotation"]["source_file_name"] = source
        record["quotation"]["source_sheet_name"] = source
        record["quotation_items"][0]["description"] = source
        self._load([record])
        self.assertEqual(self._query("SELECT source_file_name, source_sheet_name FROM {schema}.quotations"), [(source, source)])
        self.assertEqual(self._query("SELECT description FROM {schema}.quotation_items"), [(source,)])
        self.assertEqual(self._count("customers"), 1)
        self.assertEqual(self._count("quotations"), 1)

    def test_component_failure_rolls_back_entire_batch_including_masters(self):
        self._reject_component_trigger()
        before = self._snapshot()
        failed = _record(2, customer="Later batch customer", documents=True)
        failed["kit_documents"][1]["groups"][0]["components"][0]["part_name"] = "FAIL_COMPONENT"
        with self.assertRaises(LoadError):
            self._load([_record(customer="First batch customer"), failed])
        self.assertEqual(self._snapshot(), before)

    def test_failed_replace_restores_original_graph_and_master_rows(self):
        self._load([_record(documents=True)])
        self._reject_component_trigger()
        before = self._snapshot()
        replacement = _record(customer="Replacement customer", documents=True)
        replacement["kit_documents"][0]["groups"][0]["components"][0]["part_name"] = "FAIL_COMPONENT"
        with self.assertRaises(LoadError):
            self._load([replacement], existing_policy="replace")
        self.assertEqual(self._snapshot(), before)

    def test_validation_blocks_writes_and_explicit_skip_invalid_loads_only_valid(self):
        invalid = _record(2)
        invalid["quotation"]["currency"] = None
        with self.assertRaises(ValidationError):
            self._load([_record(), invalid])
        self.assertTrue(all(count == 0 for count in self._counts().values()))
        result = self._load([_record(), invalid], skip_invalid=True)
        self.assertEqual((result["inserted"], result["rejected"]), (1, 1))
        self.assertEqual(self._count("quotations"), 1)

    def test_warning_gate_requires_explicit_allow_warnings(self):
        record = _record()
        record["issues"].append({"code": "source_note", "field": "raw",
                                 "message": "Synthetic source note", "severity": "warning"})
        with self.assertRaises(ValidationError):
            self._load([record])
        self.assertEqual(self._count("customers"), 0)
        self.assertEqual(self._load([record], allow_warnings=True)["inserted"], 1)

    def test_initialize_is_idempotent_and_preserves_data(self):
        self._load([_record(documents=True)])
        before = self._snapshot()
        initialize_schema(self.connection, schema=self.schema)
        self.assertEqual(self._snapshot(), before)

    def test_partial_schema_initialization_refuses_and_preserves_existing_objects(self):
        self._drop_schema()
        self.connection.execute(sql.SQL("CREATE SCHEMA {}").format(sql.Identifier(self.schema)))
        self.connection.execute(sql.SQL("CREATE TABLE {}.existing_data (payload text)").format(sql.Identifier(self.schema)))
        self.connection.execute(sql.SQL("INSERT INTO {}.existing_data VALUES (%s)").format(sql.Identifier(self.schema)), ("preserve me",))
        with self.assertRaisesRegex(LoadError, "empty schema"):
            initialize_schema(self.connection, schema=self.schema)
        self.assertEqual(self._query("SELECT payload FROM {schema}.existing_data"), [("preserve me",)])
        self.assertEqual(self._query("SELECT count(*) FROM information_schema.tables WHERE table_schema = %s", (self.schema,)), [(1,)])

    def test_missing_unique_reference_index_prevents_loading_or_reinitializing(self):
        self.connection.execute(sql.SQL("ALTER TABLE {}.quotations DROP CONSTRAINT uq_quotations_ref_no").format(sql.Identifier(self.schema)))
        with self.assertRaisesRegex(LoadError, "unique quotation_ref_no"):
            self._load([_record()])
        with self.assertRaisesRegex(LoadError, "unique quotation_ref_no"):
            initialize_schema(self.connection, schema=self.schema)
        self.assertEqual(self._count("customers"), 0)
        self.assertEqual(self._count("quotations"), 0)

    def _assert_incompatible_schema_preserved(self, message):
        before = self._snapshot()
        with self.assertRaisesRegex(LoadError, message):
            self._load([_record()])
        with self.assertRaisesRegex(LoadError, message):
            initialize_schema(self.connection, schema=self.schema)
        self.assertEqual(self._snapshot(), before)

    def test_missing_component_foreign_key_refuses_load_and_reinitialization(self):
        self.connection.execute(sql.SQL("""
            ALTER TABLE {}.kit_components DROP CONSTRAINT fk_kit_components_component_group
        """).format(sql.Identifier(self.schema)))
        self._assert_incompatible_schema_preserved("missing or incompatible foreign key")

    def test_missing_required_not_null_refuses_load_and_reinitialization(self):
        self.connection.execute(sql.SQL("""
            ALTER TABLE {}.customers ALTER COLUMN customer_name DROP NOT NULL
        """).format(sql.Identifier(self.schema)))
        self._assert_incompatible_schema_preserved("must be NOT NULL")

    def test_incompatible_money_precision_refuses_load_and_reinitialization(self):
        self.connection.execute(sql.SQL("""
            ALTER TABLE {}.quotation_items ALTER COLUMN unit_price TYPE numeric(15, 2)
        """).format(sql.Identifier(self.schema)))
        self._assert_incompatible_schema_preserved("incompatible type")

    def test_missing_primary_key_refuses_load_and_reinitialization(self):
        self.connection.execute(sql.SQL("""
            ALTER TABLE {}.kit_components DROP CONSTRAINT kit_components_pkey
        """).format(sql.Identifier(self.schema)))
        self._assert_incompatible_schema_preserved("incompatible primary key")

    def test_non_generated_primary_key_refuses_load_and_reinitialization(self):
        self.connection.execute(sql.SQL("""
            ALTER TABLE {}.kit_components ALTER COLUMN kit_component_id DROP IDENTITY
        """).format(sql.Identifier(self.schema)))
        self._assert_incompatible_schema_preserved("must be generated")

    def test_dict_row_connection_is_supported_without_changing_caller_row_factory(self):
        with psycopg.connect(self.dsn, autocommit=True, connect_timeout=5,
                             row_factory=dict_row) as connection:
            initialize_schema(connection, schema=self.schema)
            result = load_records(connection, [_record(documents=True), _record(2)], schema=self.schema)
            self.assertEqual(result["inserted"], 2)
            self.assertIs(connection.row_factory, dict_row)
            row = connection.execute(sql.SQL("SELECT count(*) AS total FROM {}.quotations").format(
                sql.Identifier(self.schema))).fetchone()
            self.assertEqual(row, {"total": 2})
        self.assertEqual(self._count("kit_components"), 2)

    def test_serializable_session_load_uses_read_committed_without_changing_caller_setting(self):
        self.connection.execute(sql.SQL("CREATE TABLE {}.observed_isolation (value text)").format(
            sql.Identifier(self.schema)))
        self.connection.execute(sql.SQL("""
            CREATE FUNCTION {schema}.capture_isolation() RETURNS trigger
            LANGUAGE plpgsql AS $body$
            BEGIN
                INSERT INTO {schema}.observed_isolation
                    VALUES (current_setting('transaction_isolation'));
                RETURN NEW;
            END
            $body$
        """).format(schema=sql.Identifier(self.schema)))
        self.connection.execute(sql.SQL("""
            CREATE TRIGGER capture_isolation BEFORE INSERT ON {schema}.quotations
            FOR EACH ROW EXECUTE FUNCTION {schema}.capture_isolation()
        """).format(schema=sql.Identifier(self.schema)))
        with psycopg.connect(self.dsn, autocommit=True, connect_timeout=5) as connection:
            connection.isolation_level = IsolationLevel.SERIALIZABLE
            initialize_schema(connection, schema=self.schema)
            result = load_records(connection, [_record()], schema=self.schema)
            self.assertEqual(result["inserted"], 1)
            self.assertEqual(connection.isolation_level, IsolationLevel.SERIALIZABLE)
            self.assertEqual(connection.info.transaction_status, TransactionStatus.IDLE)
        self.assertEqual(self._query("SELECT value FROM {schema}.observed_isolation"), [("read committed",)])

    def test_replace_refuses_cross_quotation_item_references_without_changing_data(self):
        self._load([_record(), _record(2, documents=True)])
        item_id = self._query("""
            SELECT i.quotation_item_id FROM {schema}.quotation_items i
            JOIN {schema}.quotations q ON q.quotation_id = i.quotation_id
            WHERE q.quotation_ref_no = 'KP-M-Q-26-001'
        """)[0][0]
        self.connection.execute(sql.SQL("UPDATE {}.kit_documents SET kit_quotation_item_id = %s").format(
            sql.Identifier(self.schema)), (item_id,))
        before = self._snapshot()
        with self.assertRaisesRegex(LoadError, "referenced by another quotation"):
            self._load([_record(customer="Replacement customer")], existing_policy="replace")
        self.assertEqual(self._snapshot(), before)

    def test_caller_transaction_remains_open_and_is_never_committed_by_loader(self):
        with self.connection.transaction():
            self.connection.execute("SELECT 1")
            self.assertEqual(self.connection.info.transaction_status, TransactionStatus.INTRANS)
            with self.assertRaisesRegex(LoadError, "idle connection"):
                self._load([_record()])
            with self.assertRaisesRegex(LoadError, "idle connection"):
                initialize_schema(self.connection, schema=self.schema)
            self.assertEqual(self.connection.info.transaction_status, TransactionStatus.INTRANS)
            self.assertEqual(self._count("quotations"), 0)
        self.assertEqual(self.connection.info.transaction_status, TransactionStatus.IDLE)

    def test_workbook_pipeline_then_json_cli_repeat_load(self):
        self._drop_schema()
        with TemporaryDirectory() as directory:
            directory = Path(directory)
            workbook, first_report = directory / "synthetic.xlsx", directory / "first.json"
            _write_workbook(workbook)
            status, output = self._run_cli(pipeline_main, [
                str(workbook), "--sheet", "quote", "--kit-sheet", "kit", "--cost-sheet", "cost",
                "--cost-currency", "KRW", "--init-schema",
                *self._database_cli_arguments(first_report),
            ])
            self.assertEqual(status, 0, output)
            first = json.loads(first_report.read_text(encoding="utf-8"))
            self.assertEqual(first["status"], "committed")
            self.assertEqual(first["load"]["inserted"], 1)
            self.assertEqual(len(first["transformed"]["records"][0]["kit_documents"]), 2)
            self.assertEqual(self._count("kit_components"), 2)
            self.assertEqual(self._query("SELECT part_no FROM {schema}.kit_components ORDER BY kit_component_id"), [("000019-A",)] * 2)
            self.assertNotIn(self.dsn, first_report.read_text(encoding="utf-8"))
            before = self._snapshot()
            staging, repeat_report = directory / "staging.json", directory / "repeat.json"
            staging.write_text(json.dumps(first["transformed"]), encoding="utf-8")
            status, output = self._run_cli(json_load_main, [
                str(staging), *self._database_cli_arguments(repeat_report),
            ])
            self.assertEqual(status, 0, output)
            repeated = json.loads(repeat_report.read_text(encoding="utf-8"))
            self.assertEqual(repeated["status"], "committed")
            self.assertEqual((repeated["load"]["inserted"], repeated["load"]["skipped"]), (0, 1))
            self.assertEqual(repeated["load"]["records"][0]["quotation_id"], first["load"]["records"][0]["quotation_id"])
            self.assertEqual(self._snapshot(), before)

    def test_workbook_pipeline_validation_failure_preserves_existing_database(self):
        self._load([_record(documents=True)])
        before = self._snapshot()
        with TemporaryDirectory() as directory:
            directory = Path(directory)
            workbook, report_path = directory / "invalid.xlsx", directory / "invalid.json"
            _write_workbook(workbook, valid_customer=False)
            status, output = self._run_cli(pipeline_main, [
                str(workbook), "--sheet", "quote", *self._database_cli_arguments(report_path),
            ])
            self.assertEqual(status, 1, output)
            report = json.loads(report_path.read_text(encoding="utf-8"))
            self.assertEqual(report["status"], "validation_failed")
            self.assertEqual((report["load"]["inserted"], report["load"]["rejected"]), (0, 1))
            self.assertEqual(self._snapshot(), before)

    def test_json_cli_component_failure_reports_failure_and_preserves_database(self):
        self._load([_record(documents=True)])
        self._reject_component_trigger()
        before = self._snapshot()
        failed = _record(3, customer="Later CLI customer", documents=True)
        failed["kit_documents"][0]["groups"][0]["components"][0]["part_name"] = "FAIL_COMPONENT"
        with TemporaryDirectory() as directory:
            directory = Path(directory)
            staging, report_path = directory / "batch.json", directory / "failure.json"
            staging.write_text(json.dumps(to_jsonable({
                "source_file_name": "synthetic.xlsx", "skipped_sheets": [],
                "records": [_record(2, customer="First CLI customer"), failed],
            })), encoding="utf-8")
            status, output = self._run_cli(json_load_main, [
                str(staging), *self._database_cli_arguments(report_path),
            ])
            self.assertEqual(status, 2, output)
            report = json.loads(report_path.read_text(encoding="utf-8"))
            self.assertEqual(report["status"], "failed")
            self.assertEqual(report["stage"], "database")
            self.assertNotIn("FAIL_COMPONENT", report["error"])
            self.assertNotIn("FAIL_COMPONENT", output)
            self.assertNotIn(self.dsn, report_path.read_text(encoding="utf-8"))
            self.assertEqual(self._snapshot(), before)

    def test_concurrent_same_reference_creates_one_complete_graph(self):
        barrier = Barrier(2)
        record = _record(documents=True)

        def concurrent_load():
            with psycopg.connect(self.dsn, autocommit=True, connect_timeout=5,
                                 options="-c statement_timeout=10000 -c lock_timeout=10000") as connection:
                connection.isolation_level = IsolationLevel.REPEATABLE_READ
                barrier.wait(timeout=10)
                result = load_records(connection, [deepcopy(record)], schema=self.schema)
                self.assertEqual(connection.isolation_level, IsolationLevel.REPEATABLE_READ)
                return result

        with ThreadPoolExecutor(max_workers=2) as pool:
            results = list(pool.map(lambda _: concurrent_load(), range(2)))
        self.assertEqual(sorted(result["inserted"] for result in results), [0, 1])
        self.assertEqual(sorted(result["skipped"] for result in results), [0, 1])
        self.assertEqual(len({result["records"][0]["quotation_id"] for result in results}), 1)
        self.assertEqual(self._counts(), {
            "customers": 1, "contacts": 1, "vessels": 1, "quotations": 1,
            "quotation_vessels": 1, "quotation_items": 1, "kit_documents": 2,
            "kit_component_groups": 2, "kit_components": 2, "service_details": 0,
        })


if __name__ == "__main__":
    unittest.main()
