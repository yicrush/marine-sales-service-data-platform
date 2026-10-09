"""Regression coverage for supported M table layouts using invented sources."""

import unittest

from openpyxl import Workbook

from etl.extract.m_quotation import extract_m_items, extract_m_quotation, extract_m_summary


class MQuotationExtractTests(unittest.TestCase):
    def setUp(self):
        self.book = Workbook()
        self.addCleanup(self.book.close)
        self.sheet = self.book.active
        self.sheet.title = "Invented quote"
        for address, value in {
            "A1": "REF. NO. : KP-M-Q-26-001", "A2": "TO : Invented Marine",
            "A13": "Synthetic spares", "B15": "No.", "C15": "Description",
            "H15": "Qty", "I15": "U/PRICE", "J15": "AMOUNT", "K15": "Remark",
        }.items():
            self.sheet[address] = value

    def item(self, row, number, description="Invented part", quantity=2, unit_price=10, amount=20):
        for column, value in {2: number, 3: description, 8: quantity, 9: unit_price, 10: amount}.items():
            self.sheet.cell(row, column, value)

    def total(self, row, label="TOTAL(USD)", amount=20):
        self.sheet.cell(row, 8, label)
        self.sheet.cell(row, 10, amount)

    def test_blank_line_numbers_keep_every_independently_priced_row(self):
        self.item(16, None, "First invented part")
        self.item(17, None, "Second invented part", quantity=1, unit_price=7, amount=7)
        self.sheet["H18"] = "DISCOUNT"
        self.sheet["I18"] = 0.1
        self.sheet["J18"] = 2.7
        self.total(19, amount=24.3)
        self.item(20, 99, "Not an item after the final total")
        items = extract_m_items(self.sheet)
        self.assertEqual(len(items), 2)
        self.assertEqual([item["line_no"] for item in items], [1, 2])
        self.assertEqual([item["source_line_no"] for item in items], [None, None])
        self.assertEqual([item["source_row"] for item in items], [16, 17])
        self.assertEqual([item["amount"] for item in items], [20, 7])
        self.assertEqual(items[1]["description_raw"], "Second invented part")
        self.assertEqual(extract_m_summary(self.sheet)["quoted_amount"], 24.3)

    def test_equipment_headings_and_restarted_numbers_preserve_context(self):
        self.sheet["B16"] = "Main pump MODEL-A"
        self.item(17, 1, "Main seal")
        self.sheet["B19"] = "Servo pump MODEL-B"
        self.item(20, 1, "Servo seal")
        self.total(21, amount=40)
        items = extract_m_items(self.sheet)
        self.assertEqual([item["line_no"] for item in items], [1, 2])
        self.assertEqual([item["source_line_no"] for item in items], [1, 1])
        self.assertEqual([item["equipment_heading_raw"] for item in items], ["Main pump MODEL-A", "Servo pump MODEL-B"])
        self.assertEqual([item["source_row"] for item in items], [17, 20])

    def test_blank_separator_does_not_stop_or_reset_equipment_context(self):
        self.sheet["A16"] = "Equipment MODEL-A"
        self.item(17, 2)
        self.item(19, 7)
        self.total(20, amount=40)
        items = extract_m_items(self.sheet)
        self.assertEqual([item["line_no"] for item in items], [2, 7])
        self.assertEqual([item["source_line_no"] for item in items], [2, 7])
        self.assertEqual([item["equipment_heading_raw"] for item in items], ["Equipment MODEL-A"] * 2)

    def test_subtotals_and_repeated_headers_keep_later_equipment_blocks(self):
        self.sheet["A16"] = "Equipment MODEL-A"
        self.item(17, 1)
        self.sheet["H18"] = "SUB TOTAL(A)"
        self.sheet["J18"] = 20
        self.sheet["A19"] = "Equipment MODEL-B"
        for column, value in {2: "No.", 3: "Description", 8: "Qty", 9: "U/PRICE", 10: "AMOUNT"}.items():
            self.sheet.cell(21, column, value)
        self.item(22, 1)
        self.sheet["H23"] = "SUBTOTAL(B)"
        self.sheet["J23"] = 20
        self.sheet["H24"] = "SUB TOTAL (A) + (B)"
        self.sheet["J24"] = 40
        self.total(25, label="TOTAL / vessel", amount=40)
        record = extract_m_quotation(self.sheet)
        self.assertEqual(len(record["items"]), 2)
        self.assertEqual([item["line_no"] for item in record["items"]], [1, 2])
        self.assertEqual(record["items"][1]["equipment_heading_raw"], "Equipment MODEL-B")
        self.assertEqual(record["summary"]["quoted_amount_label_raw"], "TOTAL / vessel")
        self.assertEqual(record["summary"]["quoted_amount"], 40)
        subtotal_rows = [row for row in record["unpriced_rows_raw"] if row["source_row"] in (18, 23, 24)]
        self.assertEqual(len(subtotal_rows), 3)
        self.assertEqual(subtotal_rows[0]["cells_raw"][7], "SUB TOTAL(A)")

    def test_subtotal_is_never_relabelled_as_final_total(self):
        self.item(16, 1)
        self.sheet["H17"] = "SUB TOTAL(A)"
        self.sheet["J17"] = 20
        self.sheet["A18"] = "TERMS AND CONDITIONS:"
        summary = extract_m_summary(self.sheet)
        self.assertIsNone(summary["quoted_amount"])
        self.assertIsNone(summary["quoted_amount_label_raw"])

    def test_total_in_item_name_does_not_end_extraction(self):
        self.item(16, None, "Total repair kit")
        self.item(17, None, "Discount control valve")
        self.total(18, amount=40)
        items = extract_m_items(self.sheet)
        self.assertEqual([item["description_raw"] for item in items], ["Total repair kit", "Discount control valve"])

    def test_footer_stops_before_later_numeric_or_priced_rows(self):
        for label in ("TERMS AND CONDITIONS:", "LEAD TIME : 2 weeks", "PAYMENT TERM : Example terms"):
            with self.subTest(label=label):
                self.sheet["A18"] = label
                self.item(16, 1)
                self.item(19, 2, "Footer content must not become an item")
                self.assertEqual(len(extract_m_items(self.sheet)), 1)

    def test_invalid_numbers_are_not_repaired_when_other_rows_need_positions(self):
        for row, number in enumerate((0, -1, 1.5, True), 16):
            self.item(row, number)
        self.item(20, None)
        self.total(21, amount=100)
        items = extract_m_items(self.sheet)
        self.assertEqual([item["line_no"] for item in items], [0, -1, 1.5, True, 5])
        self.assertIs(items[3]["line_no"], True)
        self.assertEqual([item["source_line_no"] for item in items], [0, -1, 1.5, True, None])

    def test_unexpected_priced_text_number_is_preserved_for_validation(self):
        self.item(16, "unexpected-label")
        self.total(17)
        item = extract_m_items(self.sheet)[0]
        self.assertEqual(item["line_no"], "unexpected-label")
        self.assertEqual(item["source_line_no"], "unexpected-label")

    def test_no_label_on_priced_row_is_not_mistaken_for_a_repeated_header(self):
        self.item(16, "No.")
        self.total(17)
        item = extract_m_items(self.sheet)[0]
        self.assertEqual(item["source_line_no"], "No.")
        self.assertEqual(item["line_no"], "No.")

    def test_printed_numeric_strings_and_letter_labels_remain_in_raw(self):
        self.item(16, "001")
        self.item(17, "A")
        self.total(18, amount=40)
        items = extract_m_items(self.sheet)
        self.assertEqual([item["line_no"] for item in items], [1, 2])
        self.assertEqual([item["source_line_no"] for item in items], ["001", "A"])

    def test_numbered_unpriced_rows_are_retained_without_price_fill(self):
        self.item(16, 1, "Original unpriced description", quantity=None, unit_price=None, amount=None)
        self.item(17, 3, None, quantity=None, unit_price=None, amount=None)
        self.item(18, 4)
        self.total(19)
        items = extract_m_items(self.sheet)
        self.assertEqual(len(items), 3)
        self.assertEqual([item["line_no"] for item in items], [1, 3, 4])
        self.assertIsNone(items[0]["unit_price"])
        self.assertIsNone(items[1]["description_raw"])
        self.assertIsNone(items[1]["amount"])

    def test_ambiguous_description_only_row_stays_raw_without_becoming_heading(self):
        self.item(16, 1)
        self.sheet["C17"] = "Additional unpriced source text\nOriginal line break"
        self.item(18, 2)
        self.total(19, amount=40)
        record = extract_m_quotation(self.sheet)
        self.assertEqual(len(record["items"]), 2)
        self.assertIsNone(record["items"][1]["equipment_heading_raw"])
        source_row = next(row for row in record["unpriced_rows_raw"] if row["source_row"] == 17)
        self.assertEqual(source_row["cells_raw"][2], "Additional unpriced source text\nOriginal line break")

    def test_zero_prices_are_source_values_and_worksheet_is_not_modified(self):
        self.item(16, None, quantity=0, unit_price=0, amount=0)
        self.total(17, amount=0)
        before = tuple(self.sheet.iter_rows(values_only=True))
        item = extract_m_items(self.sheet)[0]
        self.assertEqual((item["quantity"], item["unit_price"], item["amount"]), (0, 0, 0))
        self.assertEqual(tuple(self.sheet.iter_rows(values_only=True)), before)

    def test_missing_exact_header_does_not_guess_an_item_table(self):
        self.sheet["B15"] = "Unrelated heading"
        self.item(16, 1)
        self.assertEqual(extract_m_items(self.sheet), [])
        self.assertEqual(extract_m_quotation(self.sheet)["unpriced_rows_raw"], [])


if __name__ == "__main__":
    unittest.main()
