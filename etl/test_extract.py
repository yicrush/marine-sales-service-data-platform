from pprint import pprint

from openpyxl import load_workbook

from etl.extract.m_quotation import extract_m_quotation
from etl.extract.o_quotation import extract_o_quotation
from etl.extract.kit_document import (extract_kit_detail, extract_cost_sheet,)


# =========================
# Test settings
# =========================

M_FILE_PATH = "data/raw/KP-M-Q-26-000 FROM START HERE.xlsx"
M_SHEET_NAME = "quote25-002-1"

O_FILE_PATH = "data/raw/KP-O-Q-26-000.xlsx"
O_SHEET_NAME = " "
KIT_SHEET_NAME = " "
COST_SHEET_NAME = " "


# =========================
# M quotation test
# =========================

m_workbook = load_workbook(
    M_FILE_PATH,
    data_only=True
)

m_sheet = m_workbook[M_SHEET_NAME]

m_data = extract_m_quotation(m_sheet)

print("\n=== M QUOTATION ===")
pprint(m_data, width=200)


# =========================
# O / KIT / COST workbook
# =========================

o_workbook = load_workbook(
    O_FILE_PATH,
    data_only=True
)


# O quotation
o_sheet = o_workbook[O_SHEET_NAME]

o_data = extract_o_quotation(o_sheet)

print("\n=== O QUOTATION ===")
pprint(o_data, width=200)


# KIT detail
kit_sheet = o_workbook[KIT_SHEET_NAME]

kit_data = extract_kit_detail(kit_sheet)

print("\n=== KIT DETAIL ===")
pprint(kit_data, width=200)


# Cost sheet
cost_sheet = o_workbook[COST_SHEET_NAME]

cost_data = extract_cost_sheet(cost_sheet)

print("\n=== COST SHEET ===")
pprint(cost_data, width=200)