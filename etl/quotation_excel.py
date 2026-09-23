from openpyxl import load_workbook


file_path = "data/raw/KP-M-Q-26-000 FROM START HERE.xlsx"

workbook = load_workbook(file_path, data_only=True)

print(f"Number of sheets: {len(workbook.sheetnames)}")
print(workbook.sheetnames[:10])

sheet_name = workbook.sheetnames[101]
sheet = workbook[sheet_name]

print(f"Sheet: {sheet_name}")
"""
for row in sheet.iter_rows():
    for cell in row:
        if cell.value is not None:
            print(f"{cell.coordinate}: {cell.value}")
"""
def find_value_by_label(sheet, label):
    for row in sheet.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and label in cell.value:
                return cell.value

    return None

print(find_value_by_label(sheet, "REF. NO."))
print(find_value_by_label(sheet, "VESSEL NAME"))
print(find_value_by_label(sheet, "DATE :"))