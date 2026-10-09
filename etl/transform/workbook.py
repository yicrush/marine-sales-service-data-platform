"""Read-only workbook orchestration around the existing Extract functions.

Supporting KIT/cost documents require an explicit association with one quotation.
No workbook is saved, and no matching is inferred from sheet order or names.
"""

from dataclasses import dataclass
from pathlib import Path
import re

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from etl.extract.common import detect_sheet_type, is_number
from etl.extract.kit_document import extract_cost_sheet, extract_kit_detail
from etl.extract.m_quotation import extract_m_quotation
from etl.extract.o_quotation import extract_o_quotation

from .kit_document import transform_kit_document
from .quotation import transform_quotation


@dataclass(frozen=True)
class _Cell:
    value: object
    row: int
    column: int
    number_format: str = "General"
    data_type: str = "n"

    @property
    def coordinate(self):
        return f"{get_column_letter(self.column)}{self.row}"


class _SheetSnapshot:
    """Provide Extract's cell-access interface without repeated ZIP rescans.

    openpyxl read-only worksheets stream rows. Extract calls ``cell`` repeatedly,
    so one in-memory snapshot makes those existing calls inexpensive while the
    source workbook stays read-only.
    """

    def __init__(self, source):
        self.title = source.title
        self._rows = []
        for row_number, row in enumerate(source.iter_rows(), 1):
            # Excel frequently preserves formatting far beyond the actual data.
            # Keep interior blank cells/formats but do not retain padded columns.
            last_value = max((index for index, cell in enumerate(row, 1)
                              if cell.value is not None), default=0)
            self._rows.append(tuple(
                _Cell(cell.value, row_number, column_number,
                      getattr(cell, "number_format", "General"),
                      getattr(cell, "data_type", "n"))
                for column_number, cell in enumerate(row[:last_value], 1)
            ))
        while self._rows and not self._rows[-1]:
            self._rows.pop()
        self.max_row = len(self._rows)
        self.max_column = max((len(row) for row in self._rows), default=0)

    def cell(self, row, column):
        if row < 1 or column < 1:
            raise ValueError("Cell coordinates must be positive")
        if row <= self.max_row and column <= len(self._rows[row - 1]):
            return self._rows[row - 1][column - 1]
        return _Cell(None, row, column)

    def __getitem__(self, row):
        return tuple(self.cell(row, column)
                     for column in range(1, self.max_column + 1))

    def iter_rows(self, min_row=1, max_row=None, min_col=1, max_col=None,
                  values_only=False):
        last_row = self.max_row if max_row is None else max_row
        last_column = self.max_column if max_col is None else max_col
        for row in range(min_row, last_row + 1):
            cells = tuple(self.cell(row, column)
                          for column in range(min_col, last_column + 1))
            yield tuple(cell.value for cell in cells) if values_only else cells


_CURRENCY_TOKEN = re.compile(
    r"(?<![A-Z])(?:US\$|USD|KRW|WON|SGD|EUR)(?![A-Z])", re.I
)
_PRICE_HEADER = re.compile(
    r"\b(?:CURRENCY|U\s*/\s*PRICE|UNIT\s*PRICE|AMOUNT|TOTAL|PRICE|COST)\b",
    re.I,
)


def _currency_codes(value, *, number_format=False):
    """Use explicit currency amounts/headers, never an unqualified dollar sign."""
    if not isinstance(value, str):
        return set()
    result = set()
    for match in _CURRENCY_TOKEN.finditer(value):
        before, after = value[:match.start()], value[match.end():]
        is_amount = bool(re.match(r"\s*[+-]?\d", after)
                         or re.search(r"\d\s*$", before))
        standalone = value.strip().upper() == match[0].upper()
        if number_format or is_amount or standalone or _PRICE_HEADER.search(value):
            code = match[0].upper()
            result.add({"US$": "USD", "WON": "KRW"}.get(code, code))
    return result


def _issue(code, field, message, severity="warning"):
    return {"code": code, "field": field, "message": message,
            "severity": severity}


def _resolve_currency(sheet, override, field):
    evidence = set()
    for row in sheet.iter_rows():
        for cell in row:
            evidence.update(_currency_codes(cell.value))
            if is_number(cell.value):
                evidence.update(_currency_codes(cell.number_format,
                                                number_format=True))
    issues = []
    if override is not None:
        normalized = str(override).strip().upper()
        if evidence and evidence != {normalized}:
            issues.append(_issue(
                "currency_override_conflict", field,
                "Explicit currency override differs from source currency evidence "
                f"({', '.join(sorted(evidence))}); override retained",
            ))
        return override, issues
    if len(evidence) == 1:
        return evidence.pop(), issues
    if len(evidence) > 1:
        issues.append(_issue(
            "ambiguous_currency", field,
            "Conflicting source currency evidence "
            f"({', '.join(sorted(evidence))}); specify an explicit currency",
            "error",
        ))
    return None, issues


def _formula_issues(formula_sheet, cached_sheet):
    missing = []
    for row in formula_sheet.iter_rows():
        for cell in row:
            if getattr(cell, "data_type", None) == "f":
                if cached_sheet.cell(cell.row, cell.column).value is None:
                    missing.append(cell.coordinate)
    if not missing:
        return []
    locations = ", ".join(missing[:12])
    if len(missing) > 12:
        locations += f", and {len(missing) - 12} more"
    return [_issue(
        "uncached_formulas", "raw",
        f"{len(missing)} formula cell(s) have no cached result ({locations}); "
        "openpyxl does not calculate formulas. Recalculate and save the source "
        "in a spreadsheet application before relying on missing amounts",
    )]


def _discount_is_fraction(sheet):
    # Mirror Extract's last matching H-column discount row, including overwrites.
    result = False
    for row in sheet.iter_rows():
        for cell in row:
            if cell.column == 8 and isinstance(cell.value, str):
                if "DISCOUNT" in cell.value.upper():
                    result = "%" in sheet.cell(cell.row, 9).number_format
    return result


def _layout_issues(sheet, kind):
    """Flag detectable fixed-column Extract limitations without changing Extract."""
    issues = []
    if kind == "O":
        return issues
    if kind == "M":
        header = next((row for row in sheet.iter_rows()
                       if any(cell.value == "No." for cell in row)), None)
        if header is None:
            return [_issue("unsupported_item_header", "raw.items",
                           "Current M Extract requires an exact 'No.' header; "
                           "the source layout is not supported")]
        columns = {str(cell.value).strip().upper(): cell.column
                   for cell in header if isinstance(cell.value, str)}
        expected = {"NO.": 2, "QTY": 8, "U/PRICE": 9, "AMOUNT": 10}
        if any(name in columns and columns[name] != column
               for name, column in expected.items()):
            issues.append(_issue(
                "shifted_extract_columns", "raw.items",
                "Source item headers differ from M Extract's fixed columns; "
                "extracted quantities/prices require review", "error",
            ))
    else:
        header = next((row for row in sheet.iter_rows()
                       if any(isinstance(cell.value, str)
                              and "PARTS NAME" in cell.value.upper()
                              for cell in row)), None)
        if header is None:
            return [_issue("unsupported_component_header", "raw.groups",
                           "No supported Parts name header was found; "
                           "review the current KIT/cost Extract layout")]
        unit_prices = [cell.column for cell in header
                       if isinstance(cell.value, str)
                       and re.search(r"(?:U\s*/\s*PRICE|UNIT\s*PRICE)",
                                     cell.value, re.I)]
        amounts = [cell.column for cell in header
                   if isinstance(cell.value, str)
                   and "AMOUNT" in cell.value.upper()]
        if ((unit_prices and unit_prices != [6])
                or (amounts and amounts != [7])):
            issues.append(_issue(
                "shifted_extract_columns", "raw.groups",
                "Source price headers differ from KIT/cost Extract's fixed "
                "F/G columns; extracted prices require review", "error",
            ))
    return issues


def _add_issues(result, issues):
    result["issues"].extend(issues)
    result["is_valid"] = not any(issue["severity"] == "error"
                                 for issue in result["issues"])
    return result


def _sheet_names(value, argument):
    if value is None:
        return None
    try:
        names = (value,) if isinstance(value, str) else tuple(value)
    except TypeError as exc:
        raise ValueError(f"{argument} must be a sheet name or an iterable of names") from exc
    if any(not isinstance(name, str) or not name for name in names):
        raise ValueError(f"{argument} must contain exact nonempty sheet names")
    return tuple(dict.fromkeys(names))


def transform_workbook(path, *, sheet_names=None, currency=None,
                       kit_sheet_names=(), cost_sheet_names=(),
                       kit_currency=None, cost_currency=None):
    """Read Excel into transformed quotation records and validation issues.

    By default only detected M/O quotations are selected. Explicit supporting
    sheets attach to exactly one selected quotation. Currency overrides may be
    supplied separately for quotations, KIT documents, and internal cost sheets.
    Neither sheet names nor bare ``$``/``\\`` imply a currency.
    """
    try:
        path = Path(path)
    except TypeError as exc:
        raise ValueError("path must identify a local workbook file") from exc
    selected = _sheet_names(sheet_names, "sheet_names")
    kits = _sheet_names(kit_sheet_names, "kit_sheet_names") or ()
    costs = _sheet_names(cost_sheet_names, "cost_sheet_names") or ()
    if selected == ():
        raise ValueError("sheet_names must select at least one quotation sheet")
    if set(kits) & set(costs):
        raise ValueError("A supporting sheet cannot be both KIT and cost")
    if not path.is_file():
        raise ValueError(f"Workbook file does not exist: {path.name}")

    cached_book = formula_book = None
    try:
        cached_book = load_workbook(path, read_only=True, data_only=True)
        formula_book = load_workbook(path, read_only=True, data_only=False)
        requested = (*selected, *kits, *costs) if selected is not None else (*kits, *costs)
        missing = [name for name in requested if name not in cached_book.sheetnames]
        if missing:
            raise ValueError("Workbook has no sheet(s): " + ", ".join(missing))
        supports = set(kits) | set(costs)
        if selected is not None and set(selected) & supports:
            raise ValueError("Quotation and supporting sheet selections overlap")

        quotation_sheets = []
        candidates = cached_book.sheetnames if selected is None else selected
        for name in candidates:
            if name in supports:
                continue
            sheet = _SheetSnapshot(cached_book[name])
            kind = detect_sheet_type(sheet)
            if kind in ("M", "O"):
                quotation_sheets.append((name, kind))
            elif selected is not None:
                raise ValueError(f"Selected sheet is not a detected M/O quotation: {name}")
        if supports and len(quotation_sheets) != 1:
            raise ValueError("Explicit KIT/cost associations require exactly one "
                             "quotation; select it with sheet_names")

        documents = []
        for names, kind, override, extract in (
            (kits, "KIT_DETAIL", kit_currency, extract_kit_detail),
            (costs, "COST_SHEET", cost_currency, extract_cost_sheet),
        ):
            for name in names:
                sheet = _SheetSnapshot(cached_book[name])
                detected_currency, issues = _resolve_currency(sheet, override,
                                                               "document.currency")
                issues += _formula_issues(formula_book[name], sheet)
                issues += _layout_issues(sheet, kind)
                document = transform_kit_document(
                    extract(sheet), source_file_name=path.name,
                    currency=detected_currency,
                )
                documents.append(_add_issues(document, issues))

        records = []
        for name, kind in quotation_sheets:
            sheet = _SheetSnapshot(cached_book[name])
            detected_currency, issues = _resolve_currency(sheet, currency,
                                                           "quotation.currency")
            issues += _formula_issues(formula_book[sheet.title], sheet)
            issues += _layout_issues(sheet, kind)
            extract = extract_m_quotation if kind == "M" else extract_o_quotation
            record = transform_quotation(
                extract(sheet), source_file_name=path.name,
                currency=detected_currency, kit_documents=documents,
                discount_rate_is_fraction=_discount_is_fraction(sheet),
            )
            records.append(_add_issues(record, issues))

        processed = supports | {name for name, _ in quotation_sheets}
        return {"source_file_name": path.name, "records": records,
                "skipped_sheets": [name for name in cached_book.sheetnames
                                   if name not in processed]}
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError(f"Unable to transform workbook {path.name}: "
                         f"{type(exc).__name__}: {exc}") from exc
    finally:
        if formula_book is not None:
            formula_book.close()
        if cached_book is not None:
            cached_book.close()
